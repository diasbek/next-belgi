#!/usr/bin/env python3
"""Parse FIPS MKTU bilingual XLSX → locale term slices + class titles."""

from __future__ import annotations

import hashlib
import json
import re
import sys
from pathlib import Path

try:
    import openpyxl
except ImportError:
    print("Install openpyxl: pip3 install openpyxl", file=sys.stderr)
    sys.exit(1)

ROOT = Path(__file__).resolve().parents[1]
XLSX = ROOT / "data/nice/raw/mktu_13_26_2lang.xlsx"
OUT = ROOT / "src/data/nice"

EN_TITLES = {
    1: "Chemicals for industry, science and photography, agriculture; unprocessed plastics; fire extinguishing; adhesives; fertilizers.",
    2: "Paints, varnishes, lacquers; preservatives against rust and wood deterioration; colorants; raw natural resins.",
    3: "Non-medicated cosmetics and toiletries; perfumery; bleaching and cleaning preparations.",
    4: "Industrial oils and greases; lubricants; fuels; candles and wicks.",
    5: "Pharmaceuticals, medical and veterinary preparations; sanitary preparations; dietetic food; disinfectants.",
    6: "Common metals and their alloys; metal building materials; metal containers; safes.",
    7: "Machines, machine tools, power tools; motors and engines; agricultural implements; vending machines.",
    8: "Hand tools and implements, hand-operated; cutlery; side arms; razors.",
    9: "Scientific, research, navigation, photographic, audiovisual, optical apparatus; computers; software.",
    10: "Surgical, medical, dental and veterinary apparatus; artificial limbs; orthopedic articles.",
    11: "Apparatus for lighting, heating, cooling, cooking, drying, ventilating, water supply and sanitary purposes.",
    12: "Vehicles; apparatus for locomotion by land, air or water.",
    13: "Firearms; ammunition and projectiles; explosives; fireworks.",
    14: "Precious metals; jewellery, precious stones; horological instruments.",
    15: "Musical instruments; music stands; conductors' batons.",
    16: "Paper and cardboard; printed matter; stationery; artists' materials; plastic sheets.",
    17: "Rubber, gutta-percha, gum, asbestos, mica; plastics in extruded form; packing and insulating materials.",
    18: "Leather and imitations; luggage and carrying bags; umbrellas; saddlery.",
    19: "Non-metallic building materials; asphalt, pitch, bitumen; non-metallic transportable buildings.",
    20: "Furniture, mirrors, picture frames; containers not of metal.",
    21: "Household or kitchen utensils and containers; combs and brushes; glassware, porcelain.",
    22: "Ropes and string; nets; tents and tarpaulins; sails; raw fibrous textile materials.",
    23: "Yarns and threads for textile use.",
    24: "Textiles and substitutes; household linen; curtains of textile or plastic.",
    25: "Clothing, footwear, headwear.",
    26: "Lace, braid and embroidery; buttons, hooks and eyes; artificial flowers; hair decorations.",
    27: "Carpets, rugs, mats and matting; linoleum; wall hangings, not of textile.",
    28: "Games, toys and playthings; video game apparatus; gymnastic and sporting articles.",
    29: "Meat, fish, poultry; preserved fruits and vegetables; dairy products; oils and fats.",
    30: "Coffee, tea, cocoa; flour; bread, pastries; confectionery; spices; sauces.",
    31: "Raw agricultural products; live animals; fresh fruits and vegetables; seeds; animal foodstuffs.",
    32: "Beers; non-alcoholic beverages; mineral waters; fruit beverages; syrups.",
    33: "Alcoholic beverages, except beers; alcoholic preparations for making beverages.",
    34: "Tobacco and tobacco substitutes; cigarettes; electronic cigarettes; smokers' articles.",
    35: "Advertising; business management, organization and administration; office functions.",
    36: "Financial, monetary and banking services; insurance; real estate affairs.",
    37: "Construction services; installation and repair services; mining extraction.",
    38: "Telecommunications services.",
    39: "Transport; packaging and storage of goods; travel arrangement.",
    40: "Treatment of materials; recycling of waste and trash; air purification; water treatment.",
    41: "Education; providing of training; entertainment; sporting and cultural activities.",
    42: "Scientific and technological services; design and development of computer hardware and software.",
    43: "Services for providing food and drink; temporary accommodation.",
    44: "Medical services; veterinary services; hygienic and beauty care; agriculture and forestry.",
    45: "Legal services; security services; personal and social services rendered by others.",
}

UZ_TITLES = {
    1: "Sanoat, fan va fotografiya, qishloq xoʻjaligi uchun kimyoviy moddalar; plastmassalar; yongʻinga qarshi vositalar; yelimlar; oʻgʻitlar.",
    2: "Boʻyoqlar, laklar; zang va yogʻochdan himoya; pigmentlar; tabiiy smolalar.",
    3: "Kosmetika va gigiyena vositalari; atir-upa; tozalash vositalari.",
    4: "Sanoat moylari; moylash materiallari; yoqilgʻilar; shamlar.",
    5: "Farmatsevtika; gigiyena; dietik oziq-ovqat; dezinfeksiya; fungitsidlar.",
    6: "Oddiy metallar; metall qurilish materiallari; seyf lar.",
    7: "Mashinalar va asboblar; motorlar; qishloq inventari; savdo avtomatlari.",
    8: "Qoʻl asboblari; pichoq-vilka; ustara.",
    9: "Ilmiy va optik asboblar; kompyuterlar; dasturiy taʼminot.",
    10: "Tibbiy va stomatologik asboblar; ortopediya; tikuv materiallari.",
    11: "Yoritish, isitish, sovutish, pishirish, ventilyatsiya, suv taʼminoti.",
    12: "Transport vositalari; quruqlik, havo va suv transporti.",
    13: "Oʻqotar qurollar; oʻq-dorilar; portlovchi moddalar; pirotexnika.",
    14: "Qimmatbaho metallar; zargarlik; soatlar.",
    15: "Musiqa asboblari.",
    16: "Qogʻoz; bosma mahsulotlar; kantselyariya; rasm ashyolari.",
    17: "Kauchuk, plastmassa; izolyatsiya; egiluvchan quvurlar.",
    18: "Charm; sumkalar; soyabonlar; egar ashyolari.",
    19: "Nommetall qurilish materiallari; asfalt; koʻchma binolar.",
    20: "Mebel; oynalar; idishlar (nommetall).",
    21: "Uy-roʻzgʻor va oshxona anjomlari; idishlar; choʻtkalar; chinni.",
    22: "Arqonlar; toʻrlar; chodirlar; yelkanlar; toʻldiruvchi materiallar.",
    23: "Ip va iplar.",
    24: "Toʻqimachilik; uy-roʻzgʻor matolari; parda.",
    25: "Kiyim-kechak; poyabzal; bosh kiyim.",
    26: "Dantellar; tugmalar; sunʼiy gullar; soch bezaklari.",
    27: "Gilamlar; linoleum; devor qoplamalari.",
    28: "Oʻyinlar va oʻyinchoqlar; sport buyumlari; Yangi yil bezaklari.",
    29: "Goʻsht, baliq; konserva meva-sabzavot; sut mahsulotlari; yogʻlar.",
    30: "Qahva, choy, kakao; un; non; shirinliklar; ziravorlar; souslar.",
    31: "Xom qishloq mahsulotlari; jonli hayvonlar; yangi meva-sabzavot; urugʻlar.",
    32: "Pivo; alkogolsiz ichimliklar; mineral suvlar; meva sharbatlari.",
    33: "Alkogolli ichimliklar (pivodan tashqari).",
    34: "Tamaki; sigaretalar; elektron sigaretalar; chekuvchi ashyolari.",
    35: "Reklama; biznes boshqaruvi; ofis xizmatlari.",
    36: "Moliyaviy va bank xizmatlari; sugʻurta; koʻchmas mulk.",
    37: "Qurilish; oʻrnatish va taʼmirlash; qazib olish.",
    38: "Telekommunikatsiya xizmatlari.",
    39: "Transport; qadoqlash va saqlash; sayohat tashkil etish.",
    40: "Materiallarni qayta ishlash; chiqindilarni qayta ishlash.",
    41: "Taʼlim; oʻqitish; koʻngilochar; sport va madaniyat.",
    42: "Ilmiy va texnologik xizmatlar; dasturiy taʼminot ishlab chiqish.",
    43: "Ovqatlanish xizmatlari; vaqtinchalik turar joy.",
    44: "Tibbiy va veterinar xizmatlar; goʻzallik; qishloq xoʻjaligi.",
    45: "Yuridik xizmatlar; xavfsizlik; shaxsiy xizmatlar.",
}


def shorten_title(text: str) -> str:
    text = re.sub(r"^Класс\s*\d+\s*", "", text.strip(), flags=re.I)
    if "Пояснения" in text:
        text = text.split("Пояснения")[0]
    text = " ".join(line.strip() for line in text.splitlines() if line.strip())
    text = re.sub(r"\s+", " ", text).strip(" ;.")
    if len(text) > 220:
        text = text[:217].rstrip() + "…"
    return text


def main() -> None:
    if not XLSX.exists():
        print(f"Missing source: {XLSX}", file=sys.stderr)
        sys.exit(1)

    OUT.mkdir(parents=True, exist_ok=True)
    wb = openpyxl.load_workbook(XLSX, read_only=True, data_only=True)

    class_titles_ru: dict[int, str] = {}
    ws2 = wb["Алфавитный по классам (рус.)"]
    for row in ws2.iter_rows(min_row=2, values_only=True):
        cls, basic, name = row[0], row[1], row[2]
        if cls is None or name is None or basic is not None:
            continue
        try:
            cn = int(cls)
        except (TypeError, ValueError):
            continue
        class_titles_ru[cn] = shorten_title(str(name))

    terms_full: list[dict] = []
    seen: set[str] = set()
    ws1 = wb["Перечень (рус., англ)"]
    for row in ws1.iter_rows(min_row=2, values_only=True):
        cls, basic, _mod, ru, en = row[0], row[1], row[2], row[3], row[4]
        if cls is None or (not ru and not en):
            continue
        try:
            cn = int(cls)
        except (TypeError, ValueError):
            continue
        if cn < 1 or cn > 45:
            continue
        ru_s = str(ru).strip() if ru else ""
        en_s = str(en).strip() if en else ""
        if not ru_s and not en_s:
            continue
        key = f"{cn}:{(ru_s or en_s).lower()}"
        if key in seen:
            continue
        seen.add(key)
        if basic is not None:
            try:
                tid = f"{cn:02d}-{int(basic)}"
            except (TypeError, ValueError):
                tid = f"{cn:02d}-{hashlib.sha1(key.encode()).hexdigest()[:10]}"
        else:
            tid = f"{cn:02d}-{hashlib.sha1(key.encode()).hexdigest()[:10]}"
        kind = "goods" if cn <= 34 else "services"
        labels = {
            "uz": ru_s or en_s,  # UZ mirrors RU until dedicated list exists
            "ru": ru_s or en_s,
            "en": en_s or ru_s,
        }
        terms_full.append(
            {
                "id": tid,
                "classNumber": cn,
                "kind": kind,
                "labels": labels,
            }
        )

    classes = []
    for cn in range(1, 46):
        kind = "goods" if cn <= 34 else "services"
        ru = class_titles_ru.get(cn) or EN_TITLES[cn]
        classes.append(
            {
                "classNumber": cn,
                "kind": kind,
                "titles": {
                    "uz": UZ_TITLES[cn],
                    "ru": ru,
                    "en": EN_TITLES[cn],
                },
            }
        )

    (OUT / "classes.json").write_text(
        json.dumps(classes, ensure_ascii=False, indent=2) + "\n",
        encoding="utf-8",
    )
    for loc in ("uz", "ru", "en"):
        slim = [
            {
                "id": t["id"],
                "classNumber": t["classNumber"],
                "kind": t["kind"],
                "label": t["labels"][loc],
            }
            for t in terms_full
        ]
        path = OUT / f"terms-{loc}.json"
        path.write_text(
            json.dumps(slim, ensure_ascii=False, separators=(",", ":")),
            encoding="utf-8",
        )
        print(f"wrote {path.relative_to(ROOT)} ({len(slim)} terms, {path.stat().st_size} bytes)")

    print(f"wrote {OUT.relative_to(ROOT)}/classes.json ({len(classes)} classes)")
    wb.close()


if __name__ == "__main__":
    main()
